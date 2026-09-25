#!/usr/bin/env python
"""Tests for `cma_saudi_xlsx.py` — run with `python tools/connectors/test_cma_saudi_xlsx.py`.

No network access is used: a tiny 16-sheet workbook is built in memory with `openpyxl`, shaped like
the three real sheets this connector reads (Table 8, 14, 15 of the live "Institutions under
supervision of CMA" file — column positions, header text and the "Total" aggregate row were copied
from a live download on 2026-09-24), then round-tripped through `wb.save()`/`BytesIO` so the real
`openpyxl.load_workbook(..., read_only=True)` parsing path is exercised, not bypassed.

What matters most, in order:

1. Per-sheet parsing lands the right value in the right firm's row.
2. The segment-priority mapping (fund_manager > asset_manager > custodian) when a firm appears in
   more than one table.
3. The **latest non-empty quarter** is taken per firm — not the first, not the last column blindly,
   and an "NA" cell is skipped rather than treated as a real (zero) figure.
4. A trailing "*" footnote marker is stripped from the name actually used, and recorded rather than
   silently dropped.
5. An empty sheet, an unparseable workbook, or a re-ordered sheet all raise `ConnectorError` rather
   than reading (or silently skipping) the wrong thing.
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))

import base  # noqa: E402
from base import ConnectorError, normalize_name  # noqa: E402
from cma_saudi_xlsx import CMASaudiXLSXConnector, SHEET_AUM, SHEET_CUSTODY, SHEET_FUNDS  # noqa: E402

# Two quarters is enough to prove "latest, not first/last blindly"; ascending, like the real file.
DEFAULT_QUARTERS = [(2025, 4), (2026, 1)]


def _fill_single_value_sheet(ws, table_number, plural, rows, quarters):
    """Table 8 / Table 15 shape: one value column per quarter."""
    ws.cell(row=9, column=4, value=f"Table({table_number}): filler title text, not read by name")
    ws.cell(row=11, column=3, value="#")
    ws.cell(row=11, column=4, value="Arabic name column")
    ws.cell(row=11, column=5,
            value=f"Capital Market Institution{'s' if plural else ''}")
    for i, (year, q) in enumerate(quarters):
        ws.cell(row=11, column=6 + i, value=f"Quarter {q} \n {year} (Million Riyal)")
    r = 12
    for idx, (name_ar, name_en, values) in enumerate(rows, start=1):
        ws.cell(row=r, column=3, value=idx)
        ws.cell(row=r, column=4, value=name_ar)
        ws.cell(row=r, column=5, value=name_en)
        for i, v in enumerate(values):
            ws.cell(row=r, column=6 + i, value=v)
        r += 1
    ws.cell(row=r, column=4, value="الإجمالي")
    ws.cell(row=r, column=5, value="Total")


def _fill_triple_value_sheet(ws, table_number, rows, quarters):
    """Table 14 shape: Public / Private / Total, three columns per quarter — only Total is read."""
    ws.cell(row=9, column=4, value=f"Table({table_number}): filler title text, not read by name")
    ws.cell(row=11, column=1, value="#")
    ws.cell(row=11, column=2, value="Arabic name column")
    ws.cell(row=11, column=3, value="Capital Market Institution")
    col = 4
    for (year, q) in quarters:
        ws.cell(row=11, column=col, value=f"Public Quarter {q}-{year}")
        ws.cell(row=11, column=col + 1, value=f"Private Quarter {q}-{year}")
        ws.cell(row=11, column=col + 2, value=f"Total Quarter {q}-{year}")
        col += 3
    r = 12
    for idx, (name_ar, name_en, totals) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=name_ar)
        ws.cell(row=r, column=3, value=name_en)
        col = 4
        for t in totals:
            ws.cell(row=r, column=col, value="NA" if t == "NA" else 1)      # Public — not read
            ws.cell(row=r, column=col + 1, value="NA" if t == "NA" else 1)  # Private — not read
            ws.cell(row=r, column=col + 2, value=t)                        # Total — the one we read
            col += 3
        r += 1
    ws.cell(row=r, column=2, value="الإجمالي")
    ws.cell(row=r, column=3, value="Total")


def build_workbook(*, aum_rows=None, funds_rows=None, custody_rows=None,
                    aum_quarters=None, funds_quarters=None, custody_quarters=None,
                    n_sheets=16, mangle_aum_title=False) -> bytes:
    """A 16-sheet workbook with real content only at positions 8, 14 and 15 — everything else is a
    blank filler sheet, exactly as the live workbook has 17 sheets and this connector reads 3."""
    if aum_rows is None:
        aum_rows = [("شركة ألفا", "Alpha Capital", [1000.0, 1200.0])]
    if funds_rows is None:
        funds_rows = [("شركة ألفا", "Alpha Capital", [1, 2])]
    if custody_rows is None:
        custody_rows = [("شركة بيتا", "Beta Custody", [500.0, 600.0])]

    wb = openpyxl.Workbook()
    wb.active.title = "filler0"
    for i in range(1, n_sheets):
        wb.create_sheet(title=f"filler{i}")

    if SHEET_AUM < n_sheets:
        _fill_single_value_sheet(wb.worksheets[SHEET_AUM], 99 if mangle_aum_title else 8, True,
                                 aum_rows, aum_quarters or DEFAULT_QUARTERS)
    if SHEET_FUNDS < n_sheets:
        _fill_triple_value_sheet(wb.worksheets[SHEET_FUNDS], 14, funds_rows,
                                 funds_quarters or DEFAULT_QUARTERS)
    if SHEET_CUSTODY < n_sheets:
        _fill_single_value_sheet(wb.worksheets[SHEET_CUSTODY], 15, False, custody_rows,
                                 custody_quarters or DEFAULT_QUARTERS)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fetch_with(workbook_bytes: bytes):
    with patch("cma_saudi_xlsx.http_get", lambda url, timeout=90: workbook_bytes):
        return CMASaudiXLSXConnector().fetch()


def by_name(entries, name):
    for e in entries:
        if e.name == name:
            return e
    raise AssertionError(f"no entry named {name!r} among {[e.name for e in entries]}")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

class TestParsing(unittest.TestCase):
    def test_a_firm_in_all_three_tables_carries_every_figure(self):
        wb = build_workbook(
            aum_rows=[("شركة ألفا", "Alpha Capital", [1000.0, 1200.0])],
            funds_rows=[("شركة ألفا", "Alpha Capital", [3, 5])],
            custody_rows=[("شركة ألفا", "Alpha Capital", [900.0, 950.0])],
        )
        entries = fetch_with(wb)
        self.assertEqual(len(entries), 1)
        e = entries[0]
        self.assertEqual(e.name, "Alpha Capital")
        self.assertEqual(e.legal_name, "شركة ألفا")
        self.assertEqual(e.raw["aum_sar_million"], 1200.0)
        self.assertEqual(e.raw["fund_count_total"], 5)
        self.assertEqual(e.raw["custody_aum_sar_million"], 950.0)

    def test_website_phone_city_licence_date_are_always_none(self):
        """This workbook publishes none of these — never invented to fill the gap."""
        entries = fetch_with(build_workbook())
        for e in entries:
            self.assertIsNone(e.website)
            self.assertIsNone(e.phone)
            self.assertIsNone(e.city)
            self.assertIsNone(e.licence_date)


# ---------------------------------------------------------------------------
# Segment priority
# ---------------------------------------------------------------------------

class TestSegmentPriority(unittest.TestCase):
    def test_funds_outranks_aum_outranks_custody(self):
        wb = build_workbook(
            aum_rows=[
                ("شركة ألفا", "Alpha Capital", [1000.0, 1200.0]),      # also in funds -> fund_manager
                ("شركة جاما", "Gamma Capital", [200.0, 250.0]),        # AUM only -> asset_manager
            ],
            funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])],
            custody_rows=[("شركة دلتا", "Delta Custody", [50.0, 60.0])],  # custody only -> custodian
        )
        entries = fetch_with(wb)
        self.assertEqual(by_name(entries, "Alpha Capital").segment, "fund_manager")
        self.assertEqual(by_name(entries, "Gamma Capital").segment, "asset_manager")
        self.assertEqual(by_name(entries, "Delta Custody").segment, "custodian")

    def test_third_party_flag_only_for_managing_segments(self):
        wb = build_workbook(
            aum_rows=[("شركة جاما", "Gamma Capital", [200.0, 250.0])],
            funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])],
            custody_rows=[("شركة دلتا", "Delta Custody", [50.0, 60.0])],
        )
        entries = fetch_with(wb)
        self.assertTrue(by_name(entries, "Gamma Capital").third_party)
        self.assertIsNone(by_name(entries, "Delta Custody").third_party)


# ---------------------------------------------------------------------------
# Latest non-empty quarter
# ---------------------------------------------------------------------------

class TestLatestQuarter(unittest.TestCase):
    def test_the_newest_quarter_with_a_real_number_wins(self):
        wb = build_workbook(aum_rows=[("شركة ألفا", "Alpha Capital", [1000.0, 1200.0])])
        e = by_name(fetch_with(wb), "Alpha Capital")
        self.assertEqual(e.raw["aum_sar_million"], 1200.0)   # Q1 2026, not Q4 2025
        self.assertEqual(e.raw["aum_quarter"], "Q1 2026")

    def test_na_in_the_latest_quarter_falls_back_to_the_previous_one(self):
        wb = build_workbook(aum_rows=[("شركة ألفا", "Alpha Capital", [1000.0, "NA"])])
        e = by_name(fetch_with(wb), "Alpha Capital")
        self.assertEqual(e.raw["aum_sar_million"], 1000.0)
        self.assertEqual(e.raw["aum_quarter"], "Q4 2025")

    def test_all_quarters_na_leaves_the_figure_none_not_zero(self):
        wb = build_workbook(aum_rows=[("شركة ألفا", "Alpha Capital", ["NA", "NA"])])
        e = by_name(fetch_with(wb), "Alpha Capital")
        self.assertIsNone(e.raw["aum_sar_million"])
        self.assertIsNone(e.raw["aum_quarter"])


# ---------------------------------------------------------------------------
# Footnote marker
# ---------------------------------------------------------------------------

class TestFootnoteMarker(unittest.TestCase):
    def test_trailing_asterisk_is_stripped_from_the_name_and_recorded(self):
        wb = build_workbook(
            aum_rows=[("شركة ألفا*", "Alpha Capital*", [1000.0, 1200.0])],
            funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])],
        )
        e = by_name(fetch_with(wb), "Alpha Capital")
        self.assertNotIn("*", e.name)
        self.assertNotIn("*", e.legal_name)
        self.assertIn("investment fund management licence", e.licence_type)


# ---------------------------------------------------------------------------
# Failure contract
# ---------------------------------------------------------------------------

class TestFailsLoudly(unittest.TestCase):
    def test_unparseable_bytes_raise(self):
        with patch("cma_saudi_xlsx.http_get", lambda url, timeout=90: b"not an excel file"):
            with self.assertRaises(ConnectorError) as ctx:
                CMASaudiXLSXConnector().fetch()
        self.assertIn("did not parse", str(ctx.exception))

    def test_too_few_sheets_raises(self):
        wb = build_workbook(n_sheets=10)  # Table 14/15 no longer exist at their expected positions
        with self.assertRaises(ConnectorError) as ctx:
            fetch_with(wb)
        self.assertIn("sheet(s)", str(ctx.exception))

    def test_a_reordered_sheet_is_caught_by_the_title_check(self):
        wb = build_workbook(mangle_aum_title=True)
        with self.assertRaises(ConnectorError) as ctx:
            fetch_with(wb)
        self.assertIn("Table (8)", str(ctx.exception))

    def test_an_empty_table_raises(self):
        wb = build_workbook(aum_rows=[])
        with self.assertRaises(ConnectorError) as ctx:
            fetch_with(wb)
        self.assertIn("Table 8", str(ctx.exception))


# ---------------------------------------------------------------------------
# Record building
# ---------------------------------------------------------------------------

class TestBuildRecord(unittest.TestCase):
    def test_aum_figure_becomes_a_readable_size_aum_string(self):
        wb = build_workbook(aum_rows=[("شركة جاما", "Gamma Capital", [200.0, 250.0])],
                            funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])])
        e = by_name(fetch_with(wb), "Gamma Capital")
        rec = CMASaudiXLSXConnector().build_record(e, "snap.json")
        self.assertIn("SAR 250.0m", rec["size"]["aum"])
        self.assertIn("Q1 2026", rec["size"]["aum"])

    def test_custody_only_firm_gets_a_labelled_custody_figure(self):
        wb = build_workbook(funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])],
                            custody_rows=[("شركة دلتا", "Delta Custody", [50.0, 60.0])])
        e = by_name(fetch_with(wb), "Delta Custody")
        rec = CMASaudiXLSXConnector().build_record(e, "snap.json")
        self.assertIn("SAR 60.0m", rec["size"]["aum"])
        self.assertIn("custody", rec["size"]["aum"])

    def test_status_is_always_nurture_since_no_licence_date_exists(self):
        wb = build_workbook()
        for e in fetch_with(wb):
            rec = CMASaudiXLSXConnector().build_record(e, "snap.json")
            self.assertEqual(rec["status"], "nurture")


# ---------------------------------------------------------------------------
# End to end against the CRM's own schema
# ---------------------------------------------------------------------------

class TestEndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self._orig = (base.REGISTERS_DIR, base.EVENTS_DIR, base.crm.COMPANIES_DIR, base.crm.RFPS_DIR)
        base.REGISTERS_DIR = self.tmp / "registers"
        base.EVENTS_DIR = self.tmp / "events"
        base.crm.COMPANIES_DIR = self.tmp / "companies"
        base.crm.RFPS_DIR = self.tmp / "rfps"
        for d in (base.REGISTERS_DIR, base.EVENTS_DIR, base.crm.COMPANIES_DIR, base.crm.RFPS_DIR):
            d.mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        base.REGISTERS_DIR, base.EVENTS_DIR, base.crm.COMPANIES_DIR, base.crm.RFPS_DIR = self._orig
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_backfill_creates_valid_records_for_every_distinct_firm(self):
        wb = build_workbook(
            aum_rows=[("شركة ألفا", "Alpha Capital", [1000.0, 1200.0]),
                     ("شركة جاما", "Gamma Capital", [200.0, 250.0])],
            funds_rows=[("شركة ألفا", "Alpha Capital", [1, 2])],
            custody_rows=[("شركة دلتا", "Delta Custody", [50.0, 60.0])],
        )
        with patch("cma_saudi_xlsx.http_get", lambda url, timeout=90: wb):
            res = CMASaudiXLSXConnector().run(dry_run=False, backfill=True)
        self.assertEqual(res["total"], 3)  # Alpha, Gamma, Delta
        files = sorted(base.crm.COMPANIES_DIR.glob("*.yaml"))
        self.assertEqual(len(files), 3)
        errors: list[str] = []
        for f in files:
            rec = base.crm.load_yaml(f)
            base.crm.validate_company(f, rec, errors, set())
            self.assertEqual(rec["regulator"], "CMA (Saudi Arabia)")
            self.assertEqual(rec["country"], "Saudi Arabia")
        self.assertEqual(errors, [])

    def test_a_registered_firm_never_needs_a_second_record(self):
        wb = build_workbook()
        with patch("cma_saudi_xlsx.http_get", lambda url, timeout=90: wb):
            CMASaudiXLSXConnector().run(dry_run=False, backfill=True)
            res2 = CMASaudiXLSXConnector().run(dry_run=False, backfill=True)
        skipped_names = [s["name"] for s in res2["skipped"]]
        self.assertIn("Alpha Capital", skipped_names)


if __name__ == "__main__":
    unittest.main(verbosity=2)
